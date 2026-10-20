"""Excel reader (.xlsx via openpyxl, .xls via xlrd) — design §8.6.

Row 1 is the header. Each non-empty sheet is one table; the full grid is kept
so shapers (anchor header, cross-tab unpivot) can reshape awkward layouts.
"""

from __future__ import annotations

import io
from datetime import date, datetime, time

from .base import Record, ReaderError, Table, TabularDocument, normalize_header


def cell_to_str(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _xlsx_sheets(data: bytes) -> list[tuple[str, list[list[str | None]]]]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ReaderError("corrupt", f"cannot open .xlsx: {exc}") from exc
    try:
        return [
            (ws.title, [[cell_to_str(v) for v in row] for row in ws.iter_rows(values_only=True)])
            for ws in wb.worksheets
        ]
    except Exception as exc:
        raise ReaderError("corrupt", f"cannot read .xlsx: {exc}") from exc
    finally:
        wb.close()


def _xls_sheets(data: bytes) -> list[tuple[str, list[list[str | None]]]]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise ReaderError("corrupt", f"cannot open .xls: {exc}") from exc
    sheets = []
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            row: list[str | None] = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).isoformat())
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append("TRUE" if cell.value else "FALSE")
                else:
                    row.append(cell_to_str(cell.value))
            rows.append(row)
        sheets.append((sheet.name, rows))
    return sheets


def _is_empty(row: list[str | None]) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def read_excel(data: bytes) -> TabularDocument:
    sheets = _xlsx_sheets(data) if data[:2] == b"PK" else _xls_sheets(data)
    tables: list[Table] = []
    for name, grid in sheets:
        if all(_is_empty(row) for row in grid):
            continue
        header = list(grid[0]) if grid else []
        while header and header[-1] is None:
            header.pop()
        columns = normalize_header(header)
        width = len(columns)
        records: list[Record] = []
        for excel_row, row in enumerate(grid[1:], start=2):
            if _is_empty(row):
                continue
            values = {columns[i]: (row[i] if i < len(row) else None) for i in range(width)}
            for i in range(width, len(row)):
                if row[i] is not None:
                    values[f"col{i + 1}"] = row[i]
            raw = "\t".join("" if v is None else v for v in row)
            records.append(Record(f"{name}!R{excel_row}", raw, values=values))
        tables.append(Table(name, ["" if h is None else h for h in header], columns, records, grid=grid))
    if not tables:
        raise ReaderError("corrupt", "workbook has no non-empty sheets")
    return TabularDocument("EXCEL", tables, None)
