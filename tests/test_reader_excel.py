from __future__ import annotations

import io

import openpyxl
import pytest

from importer.readers import parse_document
from importer.readers.base import ReaderError


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_multi_sheet_workbook():
    data = xlsx({
        "Issuer": [["Issuer ID", "State"], ["12345", "VT"]],
        "Plan": [["Plan ID", "Metal"], ["P1", "Gold"], [None, None], ["P2", 7.0]],
        "Empty": [[None]],
    })
    doc = parse_document(data, "EXCEL")
    assert doc.format == "EXCEL"
    assert [t.name for t in doc.tables] == ["Issuer", "Plan"]
    plan = doc.table("Plan")
    assert plan.columns == ["planid", "metal"]
    assert [r.row_identity for r in plan.records] == ["Plan!R2", "Plan!R4"]
    assert plan.records[1].values == {"planid": "P2", "metal": "7"}


def test_cells_beyond_header_are_kept_as_col_n():
    record = parse_document(xlsx({"S": [["a"], ["1", "extra"]]}), "EXCEL").table("S").records[0]
    assert record.values == {"a": "1", "col2": "extra"}


def test_grid_is_kept_for_shapers():
    table = parse_document(xlsx({"S": [["a", "b"], ["1", "2"]]}), "EXCEL").table("S")
    assert table.grid[0] == ["a", "b"]
    assert table.grid[1] == ["1", "2"]


def test_xls_legacy_format():
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Rates")
    ws.write(0, 0, "Plan ID")
    ws.write(0, 1, "Rate")
    ws.write(1, 0, "P1")
    ws.write(1, 1, 12.5)
    buf = io.BytesIO()
    wb.save(buf)
    table = parse_document(buf.getvalue(), "EXCEL").table("Rates")
    assert table.columns == ["planid", "rate"]
    assert table.records[0].values == {"planid": "P1", "rate": "12.5"}


def test_corrupt_xlsx():
    with pytest.raises(ReaderError) as err:
        parse_document(b"PK\x03\x04not a workbook", "EXCEL")
    assert err.value.reason_code == "corrupt"
